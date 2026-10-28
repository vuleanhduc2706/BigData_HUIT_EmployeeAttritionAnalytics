"""
Script tự động khởi tạo Slide thuyết trình PowerPoint (.pptx) chuẩn báo cáo đồ án Mã Nguồn Mở.
Đề tài: Website Bán Bánh Ngọt Online CAKEY (Laravel + MySQL)
Trường: Đại học Công Thương TP.HCM (HUIT)
"""
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

def build_presentation():
    prs = Presentation()
    # Kích thước 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    COLOR_BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    COLOR_PRIMARY = RGBColor(30, 58, 138)      # Blue 900
    COLOR_ACCENT = RGBColor(225, 29, 72)       # Rose 600
    COLOR_SECONDARY = RGBColor(109, 40, 217)   # Purple 700
    COLOR_TEXT_DARK = RGBColor(30, 41, 59)     # Slate 800
    COLOR_TEXT_LIGHT = RGBColor(248, 250, 252) # Slate 50
    COLOR_MUTED = RGBColor(100, 116, 139)      # Slate 500
    COLOR_CARD_BG = RGBColor(241, 245, 249)    # Slate 100
    COLOR_WHITE = RGBColor(255, 255, 255)

    blank_layout = prs.slide_layouts[6] # Blank slide

    def add_header(slide, title_text, category_text="BÁO CÁO ĐỒ ÁN MÃ NGUỒN MỞ"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_before = Pt(4)

        # Line divider
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_ACCENT
        line.line.color.rgb = COLOR_ACCENT

    # =========================================================================
    # SLIDE 1: TRANG BÌA (TITLE SLIDE)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG_DARK
    bg1.line.fill.background()

    # Title box
    t1_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(10.9), Inches(5.0))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True

    p_u = tf1.paragraphs[0]
    p_u.text = "TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH (HUIT)\nKHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN CÔNG NGHỆ PHẦN MỀM"
    p_u.font.size = Pt(13)
    p_u.font.bold = True
    p_u.font.color.rgb = RGBColor(148, 163, 184)

    p_m = tf1.add_paragraph()
    p_m.text = "\nBÁO CÁO ĐỒ ÁN HỌC PHẦN: MÃ NGUỒN MỞ"
    p_m.font.size = Pt(16)
    p_m.font.bold = True
    p_m.font.color.rgb = COLOR_ACCENT

    p_main = tf1.add_paragraph()
    p_main.text = "XÂY DỰNG WEBSITE BÁN BÁNH NGỌT CAKEY"
    p_main.font.size = Pt(32)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_WHITE
    p_main.space_before = Pt(8)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Hệ thống sàn thương mại điện tử Multi-vendor chuyên biệt ngành bánh\n(Nền tảng: Laravel 11, MySQL, Blade Template, Bootstrap 5, AJAX & VietQR API)"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = RGBColor(226, 232, 240)
    p_sub.space_before = Pt(8)

    p_info = tf1.add_paragraph()
    p_info.text = "\nSinh viên thực hiện: Vũ Lê Anh Đức (MSSV: 2001230191 - Trưởng nhóm)\nGiảng viên phụ trách học phần: Khoa CNTT - HUIT"
    p_info.font.size = Pt(13)
    p_info.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # SLIDE 2: TỔNG QUAN VÀ LÝ DO CHỌN ĐỀ TÀI
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. Tổng Quan & Lý Do Chọn Đề Tài")

    def add_card(slide, left, top, width, height, title, content, border_color=COLOR_PRIMARY):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = border_color
        
        p_c = tf.add_paragraph()
        p_c.text = content
        p_c.font.size = Pt(12)
        p_c.font.color.rgb = COLOR_TEXT_DARK
        p_c.space_before = Pt(6)

    add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Bối Cảnh & Thực Trạng Thực Tế",
             "• Ngành F&B và bánh ngọt phát triển bùng nổ, khách hàng có nhu cầu đặt bánh sinh nhật, tiệc ngọt online nhanh chóng.\n\n"
             "• Các cửa hàng bánh truyền thống gặp khó khăn trong việc quản lý đơn hàng, theo dõi thanh toán chuyển khoản và tiếp cận khách hàng trẻ.\n\n"
             "• Người tiêu dùng muốn xem trực quan hình ảnh món bánh, tra cứu thành phần, giá cả và đặt hàng giao tận nơi tiện lợi.",
             COLOR_PRIMARY)

    add_card(s2, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Mục Tiêu & Giải Pháp Đề Tài",
             "• Xây dựng Website Bán Hàng Online chuyên nghiệp (Đề tài số 3 theo gợi ý của bộ môn) với tên thương hiệu CAKEY.\n\n"
             "• Nâng cấp thành mô hình Sàn Bánh Ngọt Multi-vendor (Hỗ trợ 3 nhóm đối tượng: Khách mua, Chủ tiệm bánh, Quản trị viên).\n\n"
             "• Tối ưu trải nghiệm: Tìm kiếm tức thì AJAX Live Search, thanh toán quét mã QR VietQR tự động, xác thực OTP qua Email.\n\n"
             "• Tuân thủ 100% chuẩn kỹ thuật: Kiến trúc MVC, Eloquent ORM, Migration, Seeder, Middleware.",
             COLOR_ACCENT)

    # =========================================================================
    # SLIDE 3: CÔNG NGHỆ SỬ DỤNG & KIẾN TRÚC HỆ THỐNG
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Công Nghệ Sử Dụng & Kiến Trúc Hệ Sinh Thái")

    tech_items = [
        ("Backend Framework", "Laravel Framework 11.x\n• Kiến trúc chuẩn MVC\n• Eloquent ORM truy vấn quan hệ\n• Middleware kiểm soát phân quyền\n• Form Request Validation chặt chẽ", COLOR_PRIMARY),
        ("Cơ Sở Dữ Liệu", "MySQL Database (InnoDB)\n• 8 bảng dữ liệu quan hệ chặt chẽ\n• Ràng buộc khóa ngoại Foreign Keys\n• Hỗ trợ Transaction toàn vẹn đơn hàng\n• Migration & Seeder tự động", COLOR_SECONDARY),
        ("Frontend & Giao Diện", "Blade Template + Bootstrap 5\n• Giao diện Responsive Mobile\n• Tương tác không reload qua AJAX\n• Icon FontAwesome 6 trực quan\n• Bảng điều khiển quản trị 3D hiện đại", COLOR_ACCENT),
        ("Tích Hợp Bên Ngoài (API)", "VietQR API & Email Service\n• Tạo mã QR thanh toán động VietQR\n• Tự động khớp số tiền và cú pháp\n• Gửi mã xác thực OTP qua Email\n• Cơ chế In-App Notification tức thời", COLOR_PRIMARY)
    ]

    for idx, (title, content, col) in enumerate(tech_items):
        col_idx = idx % 2
        row_idx = idx // 2
        left = Inches(0.8 + col_idx * 6.0)
        top = Inches(1.8 + row_idx * 2.5)
        add_card(s3, left, top, Inches(5.6), Inches(2.2), title, content, col)

    # =========================================================================
    # SLIDE 4: THIẾT KẾ CƠ SỞ DỮ LIỆU (ERD)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "3. Thiết Kế Cơ Sở Dữ Liệu (CSDL) & Sơ Đồ Quan Hệ (ERD)")

    table_box = s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf4 = table_box.text_frame
    tf4.word_wrap = True

    p4_desc = tf4.paragraphs[0]
    p4_desc.text = "Hệ thống CSDL 'BanhNgot' gồm 8 bảng dữ liệu chuẩn hóa, có quan hệ ràng buộc toàn vẹn khóa ngoại (Foreign Keys):"
    p4_desc.font.size = Pt(13)
    p4_desc.font.bold = True
    p4_desc.font.color.rgb = COLOR_PRIMARY

    # Add Table
    erd_table = s4.shapes.add_table(rows=9, cols=4, left=Inches(0.8), top=Inches(2.3), width=Inches(11.7), height=Inches(4.5)).table
    erd_headers = ["Tên Bảng (Table)", "Khóa Chính (PK)", "Khóa Ngoại (FK)", "Ý Nghĩa Nghiệp Vụ & Quan Hệ"]
    for i, h in enumerate(erd_headers):
        cell = erd_table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.size = Pt(11)

    erd_data = [
        ("nguoidung (Users)", "MaND", "-", "Lưu tài khoản, mật khẩu băm Bcrypt, quyền hạn (Admin, Seller, User), trạng thái khóa"),
        ("CuaHang (Shops)", "MaCuaHang", "MaND -> nguoidung", "Thông tin tiệm bánh do Seller sở hữu (1 User NguoiBan sở hữu 1 Cửa Hàng)"),
        ("danhmuc (Categories)", "MaDanhMuc", "-", "Phân loại bánh (Bánh Mì, Bánh Ngọt, Bánh Sinh Nhật, Cupcake...)"),
        ("SanPham (Products)", "MaSanPham", "MaDanhMuc, MaCuaHang", "Món bánh: Giá bán, tồn kho, hình ảnh, mô tả, trạng thái (Đang bán, Tạm gỡ, Chờ duyệt)"),
        ("GioHang (Cart)", "id (Auto Inc)", "MaNguoiDung, MaSanPham", "Giỏ hàng tạm thời của khách hàng (lưu số lượng bánh, ngày thêm)"),
        ("DonHang (Orders)", "MaDon", "MaKhachHang -> nguoidung", "Đơn hàng: Họ tên nhận, SĐT, địa chỉ giao, PTTT (COD/Chuyển khoản), trạng thái, tổng tiền"),
        ("ChiTietDonHang (OrderItems)", "id (Auto Inc)", "MaDon, MaSanPham", "Chi tiết các món bánh trong đơn (Số lượng đặt, đơn giá tại thời điểm mua)"),
        ("ThongBao (Notifications)", "id (Auto Inc)", "MaND -> nguoidung", "Hộp thư thông báo tức thì: Đơn hàng mới, duyệt tiệm, trạng thái thanh toán")
    ]

    for r_idx, row in enumerate(erd_data, 1):
        for c_idx, val in enumerate(row):
            c = erd_table.cell(r_idx, c_idx)
            c.text = val
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            if c_idx == 0:
                p.font.bold = True
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(255, 255, 255) if r_idx % 2 == 1 else COLOR_CARD_BG

    # =========================================================================
    # SLIDE 5: SƠ ĐỒ USE CASE TỔNG QUÁT (UML USE CASE)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "4. Sơ Đồ Use Case Tổng Quát (UML Use Case)")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.8),
             "Actor 1: Khách Hàng (Customer)",
             "• Đăng ký / Đăng nhập Email OTP\n"
             "• Xem danh mục & tìm kiếm Live Search\n"
             "• Thêm / sửa / xóa giỏ hàng AJAX\n"
             "• Đặt hàng & thanh toán COD / VietQR\n"
             "• Xem lịch sử đơn hàng & Hủy đơn\n"
             "• Nộp đơn khiếu nại chuyển khoản\n"
             "• Nhận thông báo tiến độ đơn hàng\n"
             "• Nộp hồ sơ đăng ký mở tiệm bánh",
             COLOR_PRIMARY)

    add_card(s5, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.8),
             "Actor 2: Người Bán (Seller)",
             "• Đăng ký & quản lý hồ sơ tiệm bánh\n"
             "• Đăng món bánh mới (Upload ảnh)\n"
             "• Chỉnh sửa thông tin & giá bánh\n"
             "• Xóa mềm (Tạm ẩn) & Khôi phục bánh\n"
             "• Xem danh sách đơn hàng đặt tại tiệm\n"
             "• Cập nhật trạng thái đơn (Duyệt, Giao)\n"
             "• Xác nhận / Từ chối chuyển khoản\n"
             "• Nhận thông báo khi có đơn hàng mới",
             COLOR_SECONDARY)

    add_card(s5, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8),
             "Actor 3: Quản Trị Viên (Admin)",
             "• Quản lý danh sách tài khoản người dùng\n"
             "• Phê duyệt / Từ chối tiệm bánh mới\n"
             "• Khóa / Mở khóa tài khoản vi phạm\n"
             "• Kiểm duyệt toàn bộ món bánh sàn\n"
             "• Xóa cứng / Khôi phục sản phẩm\n"
             "• Quản lý và đổi trạng thái toàn bộ đơn\n"
             "• Thống kê người dùng & doanh thu sàn",
             COLOR_ACCENT)

    # =========================================================================
    # SLIDE 6: CHỨC NĂNG DÀNH CHO KHÁCH HÀNG
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "5. Chức Năng Khách Hàng (Customer / Client)")

    client_cards = [
        ("Tìm Kiếm Live Search & Lọc Bánh", 
         "• Gợi ý tức thì ngay khi gõ từ khóa bằng AJAX (không cần bấm Enter hay tải lại trang).\n"
         "• Lọc bánh linh hoạt theo danh mục (Bánh Sinh Nhật, Bánh Mì, Cupcake).\n"
         "• Phân trang danh sách bánh mượt mà với Bootstrap 5 Paginator.", COLOR_PRIMARY),
        ("Giỏ Hàng AJAX & Thanh Toán Linh Hoạt", 
         "• Thêm sản phẩm vào giỏ hàng tức thời bằng AJAX popup đẹp mắt.\n"
         "• Tự động tính toán tổng tiền, cập nhật số lượng trực tiếp.\n"
         "• Hỗ trợ 2 phương thức: Trả tiền mặt (COD) hoặc Chuyển khoản VietQR.", COLOR_ACCENT),
        ("Thanh Toán VietQR Thông Minh", 
         "• Tự động sinh mã QR Ngân hàng MBBank có sẵn số tiền và nội dung đơn hàng.\n"
         "• Khách hàng chỉ cần quét bằng app ngân hàng là khớp lệnh 100% không sợ nhầm số tiền.", COLOR_SECONDARY),
        ("Quản Lý Đơn Hàng & Khiếu Nại", 
         "• Theo dõi lộ trình đơn hàng theo thời gian thực (Đang chờ -> Đã duyệt -> Đang giao -> Đã giao).\n"
         "• Khách được quyền hủy đơn khi tiệm chưa duyệt.\n"
         "• Chức năng nộp bằng chứng khiếu nại nếu chuyển khoản bị từ chối oan.", COLOR_PRIMARY)
    ]

    for idx, (title, content, col) in enumerate(client_cards):
        col_idx = idx % 2
        row_idx = idx // 2
        left = Inches(0.8 + col_idx * 6.0)
        top = Inches(1.8 + row_idx * 2.5)
        add_card(s6, left, top, Inches(5.6), Inches(2.2), title, content, col)

    # =========================================================================
    # SLIDE 7: CHỨC NĂNG DÀNH CHO NGƯỜI BÁN (SELLER)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "6. Chức Năng Chủ Tiệm Bánh (Seller / Shop Partner)")

    seller_cards = [
        ("Nộp Hồ Sơ Mở Tiệm Bánh", 
         "• Khách hàng có thể đăng ký nâng cấp thành Seller tại /seller/register.\n"
         "• Nhập tên tiệm, địa chỉ tiệm; hệ thống tạo hồ sơ Chờ Admin duyệt và gửi thông báo tới QTV.", COLOR_PRIMARY),
        ("CRUD Quản Lý Sản Phẩm Tiệm Mình", 
         "• Đăng món bánh mới: Tên bánh, giá bán, số lượng tồn kho, danh mục, mô tả.\n"
         "• Hỗ trợ Upload ảnh trực tiếp từ máy tính lên máy chủ (lưu thư mục public/img).\n"
         "• Chỉnh sửa thông tin món bánh bất kỳ lúc nào.", COLOR_ACCENT),
        ("Cơ Chế Xóa Mềm & Khôi Phục Bánh", 
         "• Xóa mềm (Tạm ẩn): Chuyển trạng thái bánh thành Tạm gỡ để khách không thấy trên sàn.\n"
         "• Nút 'Khôi phục' giúp mở lại món bánh mà không làm mất dữ liệu lịch sử đặt hàng.", COLOR_SECONDARY),
        ("Xử Lý Đơn Hàng & Duyệt Thanh Toán", 
         "• Xem danh sách các đơn hàng khách đặt riêng cho tiệm mình.\n"
         "• Nút thao tác: Duyệt đơn, Bắt đầu giao hàng, Giao thành công.\n"
         "• Với đơn chuyển khoản: Nút 'Xác nhận thanh toán' và 'Từ chối thanh toán' kèm lý do.", COLOR_PRIMARY)
    ]

    for idx, (title, content, col) in enumerate(seller_cards):
        col_idx = idx % 2
        row_idx = idx // 2
        left = Inches(0.8 + col_idx * 6.0)
        top = Inches(1.8 + row_idx * 2.5)
        add_card(s7, left, top, Inches(5.6), Inches(2.2), title, content, col)

    # =========================================================================
    # SLIDE 8: CHỨC NĂNG QUẢN TRỊ VIÊN (ADMIN)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "7. Chức Năng Quản Trị Viên Toàn Quyền (Admin)")

    admin_cards = [
        ("Quản Lý Người Dùng & Duyệt Seller", 
         "• Xem toàn bộ tài khoản; tìm kiếm theo tên, email, SĐT, mã tài khoản.\n"
         "• Bộ lọc tab: Tất cả, Khách hàng, Chủ tiệm bánh, Chờ duyệt, Bị khóa.\n"
         "• Phê duyệt hoặc Từ chối hồ sơ xin mở tiệm bánh của đối tác.", COLOR_PRIMARY),
        ("Khóa / Mở Khóa Tài Khoản Vi Phạm", 
         "• Admin có quyền khóa tài khoản có hành vi gian lận (chuyển TrangThai = 0).\n"
         "• Tài khoản bị khóa sẽ bị hệ thống từ chối đăng nhập và hủy phiên làm việc ngay lập tức.", COLOR_ACCENT),
        ("Kiểm Duyệt Sản Phẩm Toàn Sàn", 
         "• Quản lý tất cả món bánh của mọi tiệm bánh trên toàn hệ thống.\n"
         "• Các quyền can thiệp cấp cao: Tạm gỡ bánh vi phạm, Xóa vĩnh viễn (Force Delete), Khôi phục bánh khi có khiếu nại hợp lệ.", COLOR_SECONDARY),
        ("Giám Sát Đơn Hàng Toàn Hệ Thống", 
         "• Theo dõi toàn bộ đơn hàng phát sinh trên toàn sàn.\n"
         "• Can thiệp xử lý các trường hợp tranh chấp thanh toán giữa khách hàng và tiệm bánh.\n"
         "• Cập nhật trạng thái đơn hàng khi cần thiết.", COLOR_PRIMARY)
    ]

    for idx, (title, content, col) in enumerate(admin_cards):
        col_idx = idx % 2
        row_idx = idx // 2
        left = Inches(0.8 + col_idx * 6.0)
        top = Inches(1.8 + row_idx * 2.5)
        add_card(s8, left, top, Inches(5.6), Inches(2.2), title, content, col)

    # =========================================================================
    # SLIDE 9: ĐỐI CHIẾU CÁC YÊU CẦU MỞ RỘNG (ĐIỂM CỘNG THEO ĐỀ BÀI)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "8. Các Yêu Cầu Mở Rộng Đạt Điểm Cộng Tuyệt Đối")

    ext_table = s9.shapes.add_table(rows=5, cols=3, left=Inches(0.8), top=Inches(2.0), width=Inches(11.7), height=Inches(4.8)).table
    ext_headers = ["Tiêu Chí Điểm Cộng Của Thầy", "Hiện Thực Kỹ Thuật Trong Dự Án CAKEY", "Đánh Giá Barem Điểm"]
    for i, h in enumerate(ext_headers):
        cell = ext_table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.size = Pt(12)

    ext_rows = [
        ("1. Upload File (Hình Ảnh)", 
         "• SellerController::storeProduct & updateProduct xử lý file upload.\n"
         "• Validate nghiêm ngặt định dạng ảnh (jpg, png, webp, gif) và dung lượng tối đa 5MB.\n"
         "• Đổi tên file ngẫu nhiên chống trùng lặp và lưu trữ vào public/img.", "ĐẠT ĐIỂM CỘNG (Tối đa)"),
        ("2. Tích Hợp AJAX Tương Tác", 
         "• AJAX Live Search: Gõ từ khóa gợi ý danh sách bánh ngay lập tức không reload.\n"
         "• AJAX Giỏ hàng: Thêm bánh, cập nhật số lượng mượt mà.\n"
         "• AJAX Chuông thông báo: Tải danh sách thông báo và đánh dấu đã đọc ngầm.", "ĐẠT ĐIỂM CỘNG (Tối đa)"),
        ("3. Tích Hợp API Bên Ngoài", 
         "• Tích hợp API Cổng thanh toán VietQR (https://img.vietqr.io).\n"
         "• Sinh mã QR ngân hàng MBBank tự động gắn đúng số tiền, số tài khoản và nội dung đơn.", "ĐẠT ĐIỂM CỘNG (Tối đa)"),
        ("4. Responsive Tốt Trên Mobile", 
         "• Thiết kế bằng Bootstrap 5 theo chuẩn Mobile-First (lưới col-12, col-md-6, col-lg-4).\n"
         "• Thanh điều hướng co giãn linh hoạt, bảng điều khiển tự thích ứng trên smartphone.", "ĐẠT ĐIỂM CỘNG (Tối đa)")
    ]

    for r_idx, (c1, c2, c3) in enumerate(ext_rows, 1):
        ext_table.cell(r_idx, 0).text = c1
        ext_table.cell(r_idx, 1).text = c2
        ext_table.cell(r_idx, 2).text = c3
        for c_idx in range(3):
            cell = ext_table.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_WHITE if r_idx % 2 == 1 else COLOR_CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_PRIMARY
            elif c_idx == 2:
                p.font.bold = True
                p.font.color.rgb = COLOR_ACCENT

    # =========================================================================
    # SLIDE 10: CÁC TÍNH NĂNG ĐẶC BIỆT LÀM THÊM NGOÀI ĐỀ BÀI
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "9. Các Tính Năng Đặc Biệt Làm Thêm Vượt Barem")

    bonus_cards = [
        ("1. Mô Hình Sàn Multi-vendor 3 Cấp Bậc", 
         "• Nâng cấp vượt bậc từ web bán hàng đơn giản lên mô hình Sàn thương mại điện tử đa người bán.\n"
         "• Khách hàng có thể nộp đơn xin mở tiệm bánh, Admin duyệt thì trở thành Đối tác Người bán (Seller) có không gian quản lý tiệm riêng.", COLOR_PRIMARY),
        ("2. Xác Thực Email OTP Bảo Mật 2 Bước", 
         "• Đăng ký tài khoản: Hệ thống gửi mã OTP 6 số ngẫu nhiên qua Email để kích hoạt tài khoản.\n"
         "• Quên mật khẩu: Đặt lại mật khẩu an toàn qua Email OTP, không sợ bị hack tài khoản.", COLOR_ACCENT),
        ("3. Hệ Thống Thông Báo Nội Bộ Tức Thời", 
         "• Chuông thông báo đỏ (In-App Notification) hiển thị số lượng tin chưa đọc.\n"
         "• Tự động bắn thông báo khi: Có đơn hàng mới (báo Seller), Đơn được duyệt (báo Khách), Có hồ sơ shop mới (báo Admin).", COLOR_SECONDARY),
        ("4. Quy Trình Thanh Toán 2 Chiều & Khiếu Nại", 
         "• Khách quét QR VietQR -> bấm Đã chuyển khoản -> Seller kiểm tra số dư và Xác nhận hoặc Từ chối.\n"
         "• Nếu bị từ chối oan, khách có nút Nộp khiếu nại để Admin vào cuộc giải quyết công bằng.", COLOR_PRIMARY)
    ]

    for idx, (title, content, col) in enumerate(bonus_cards):
        col_idx = idx % 2
        row_idx = idx // 2
        left = Inches(0.8 + col_idx * 6.0)
        top = Inches(1.8 + row_idx * 2.5)
        add_card(s10, left, top, Inches(5.6), Inches(2.2), title, content, col)

    # =========================================================================
    # SLIDE 11: KỊCH BẢN DEMO TRỰC TIẾP (LIVE DEMO)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "10. Kịch Bản Demo Trực Tiếp Sản Phẩm (Live Demo)")

    steps = [
        ("Bước 1: Trải nghiệm Khách Mua Hàng", 
         "• Truy cập trang chủ -> Dùng Live Search tìm 'Bánh tiramisu' -> Xem chi tiết bánh.\n"
         "• Thêm vào giỏ hàng qua AJAX -> Vào Checkout chọn thanh toán Chuyển khoản VietQR -> Hệ thống hiển thị mã QR động.\n"
         "• Đặt hàng thành công -> Chuyển sang trang Lịch sử đơn hàng để theo dõi trạng thái.", COLOR_PRIMARY),
        ("Bước 2: Nộp Hồ Sơ Đăng Ký Mở Tiệm Bánh", 
         "• Đăng nhập tài khoản khách hàng -> Vào mục 'Đăng ký bán hàng' (/seller/register).\n"
         "• Điền thông tin Tiệm bánh mới -> Nộp hồ sơ -> Hệ thống tự động gửi thông báo tới Admin.", COLOR_SECONDARY),
        ("Bước 3: Admin Phê Duyệt & Quản Trị Hệ Thống", 
         "• Đăng nhập tài khoản Admin -> Vào trang Quản lý tài khoản (/admin/users).\n"
         "• Xem danh sách Chờ duyệt -> Bấm 'Phê duyệt' -> Tài khoản lập tức được cấp quyền Seller.\n"
         "• Khám phá các chức năng: Kiểm duyệt sản phẩm, xóa mềm, khôi phục bánh.", COLOR_ACCENT),
        ("Bước 4: Seller Đăng Bánh & Xử Lý Đơn Hàng", 
         "• Đăng nhập tài khoản Seller vừa duyệt -> Vào trang Quản lý tiệm bánh (/seller/products).\n"
         "• Thêm món bánh mới có Upload file ảnh từ máy tính -> Xem danh sách đơn hàng tiệm mình (/seller/orders) -> Bấm 'Xác nhận thanh toán' và 'Bắt đầu giao bánh'.", COLOR_PRIMARY)
    ]

    for idx, (title, content, col) in enumerate(steps):
        col_idx = idx % 2
        row_idx = idx // 2
        left = Inches(0.8 + col_idx * 6.0)
        top = Inches(1.8 + row_idx * 2.5)
        add_card(s11, left, top, Inches(5.6), Inches(2.2), title, content, col)

    # =========================================================================
    # SLIDE 12: BẢNG PHÂN CÔNG CÔNG VIỆC NHÓM
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "11. Bảng Phân Công Công Việc Giữa Các Thành Viên")

    team_table = s12.shapes.add_table(rows=5, cols=4, left=Inches(0.8), top=Inches(2.0), width=Inches(11.7), height=Inches(4.8)).table
    team_headers = ["Họ Và Tên", "MSSV", "Nhiệm Vụ Được Phân Công", "Tỷ Lệ Đóng Góp"]
    for i, h in enumerate(team_headers):
        cell = team_table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.size = Pt(12)

    team_rows = [
        ("Vũ Lê Anh Đức", "2001230191", "Trưởng nhóm: Phân tích kiến trúc hệ thống, thiết kế CSDL (ERD), xây dựng Backend Laravel, bảo mật Auth/Middleware, tích hợp VietQR API & hoàn thiện báo cáo", "100% (Hoàn thành xuất sắc)"),
        ("Thành viên nhóm 2", "(Điền MSSV)", "Phát triển giao diện Blade Template, Bootstrap 5, tích hợp AJAX Live Search và giỏ hàng", "100% (Hoàn thành tốt)"),
        ("Thành viên nhóm 3", "(Điền MSSV)", "Xây dựng các chức năng phân quyền Seller (CRUD sản phẩm, upload ảnh, quản lý đơn hàng)", "100% (Hoàn thành tốt)"),
        ("Thành viên nhóm 4", "(Điền MSSV)", "Kiểm thử hệ thống (Unit Test / Feature Test), chuẩn bị dữ liệu mẫu Seeder và slide thuyết trình", "100% (Hoàn thành tốt)")
    ]

    for r_idx, row in enumerate(team_rows, 1):
        for c_idx, val in enumerate(row):
            cell = team_table.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_WHITE if r_idx % 2 == 1 else COLOR_CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_PRIMARY
            elif c_idx == 3:
                p.font.bold = True
                p.font.color.rgb = COLOR_ACCENT

    # =========================================================================
    # SLIDE 13: ĐÁNH GIÁ & HƯỚNG PHÁT TRIỂN
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "12. Đánh Giá Hệ Thống & Hướng Phát Triển")

    add_card(s13, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Ưu Điểm Cốt Lõi Đạt Được",
             "• Đáp ứng 100% chuẩn yêu cầu bắt buộc của môn học (CRUD quan hệ, Auth/Middleware, Phân trang, Tìm kiếm, Validation).\n\n"
             "• Đạt tối đa 4/4 tiêu chí điểm cộng mở rộng (Upload ảnh, AJAX, VietQR API, Mobile Responsive).\n\n"
             "• Giao diện trau chuốt, thẩm mỹ cao, trải nghiệm người dùng mượt mà.\n\n"
             "• Mã nguồn chuẩn hóa theo phong cách Laravel Clean Code, dễ bảo trì và mở rộng.",
             COLOR_PRIMARY)

    add_card(s13, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), 
             "Hướng Phát Triển Trong Tương Lai",
             "• Tích hợp hệ thống Đánh giá & Bình luận sao (Product Review & Rating) cho từng món bánh.\n\n"
             "• Kết nối API đơn vị vận chuyển bên ngoài (Giao Hàng Tiết Kiệm, Ahamove) tính phí ship tự động.\n\n"
             "• Tích hợp Cổng thanh toán trực tuyến tự động IPN Webhook (MoMo, VNPay, ZaloPay).\n\n"
             "• Xây dựng ứng dụng di động (Mobile App Flutter / React Native) đồng bộ với hệ thống API.",
             COLOR_ACCENT)

    # =========================================================================
    # SLIDE 14: LỜI CẢM ƠN (THANK YOU SLIDE)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    bg14 = s14.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg14.fill.solid()
    bg14.fill.fore_color.rgb = COLOR_BG_DARK
    bg14.line.fill.background()

    t14_box = s14.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.5))
    tf14 = t14_box.text_frame
    tf14.word_wrap = True

    p14_1 = tf14.paragraphs[0]
    p14_1.text = "XIN CHÂN THÀNH CẢM ƠN THẦY VÀ CÁC BẠN\nĐÃ LẮNG NGHE BÀI THUYẾT TRÌNH!"
    p14_1.alignment = PP_ALIGN.CENTER
    p14_1.font.size = Pt(28)
    p14_1.font.bold = True
    p14_1.font.color.rgb = COLOR_WHITE

    p14_2 = tf14.add_paragraph()
    p14_2.text = "\nCAKEY BAKERY SHOP - DỰ ÁN MÃ NGUỒN MỞ\nNhóm sinh viên sẵn sàng giải đáp các câu hỏi thảo luận (Q&A)."
    p14_2.alignment = PP_ALIGN.CENTER
    p14_2.font.size = Pt(15)
    p14_2.font.color.rgb = RGBColor(148, 163, 184)

    # Lưu tệp PowerPoint vào 2 vị trí tiện lợi:
    out_dir_1 = r"D:\Doan_BigData\BaoCao_Word"
    out_dir_2 = r"C:\laragon\www\doan_banhngot"
    os.makedirs(out_dir_1, exist_ok=True)
    os.makedirs(out_dir_2, exist_ok=True)

    pptx_path_1 = os.path.join(out_dir_1, "Slide_BaoCao_MaNguonMo_Cakey.pptx")
    pptx_path_2 = os.path.join(out_dir_2, "Slide_BaoCao_MaNguonMo_Cakey.pptx")

    prs.save(pptx_path_1)
    prs.save(pptx_path_2)

    print(f"[OK] Đã tạo thành công Slide thuyết trình PowerPoint tại:")
    print(f"  -> {pptx_path_1}")
    print(f"  -> {pptx_path_2}")

if __name__ == "__main__":
    build_presentation()
